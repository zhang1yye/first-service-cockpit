#!/usr/bin/env python3
"""
从决策系统预算周报 Excel 导入累计预算 + 累计执行到 cockpit.db
数据源: 预算周报-第一服务0525.xlsx
"""
import openpyxl, sqlite3, os

XLSX_PATH = os.path.expanduser('~/Desktop/第一服务半年冲刺/预算周报-第一服务0525.xlsx')
DB_PATH = os.path.join(os.path.dirname(__file__), '..', 'cockpit.db')

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb.active

# 读取所有华北服务中心
rows = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if row[1] != '华北地区' or row[2] is None:
        continue
    rows.append({
        'center': str(row[2]).strip(),
        'annual_budget': round((row[4] or 0) / 10000),       # 全年预算
        'cumulative_budget': round((row[5] or 0) / 10000),    # 累计预算
        'cumulative_executed': round((row[6] or 0) / 10000),  # 执行合计
    })

print(f'📊 预算周报: {len(rows)} 个华北服务中心')
total_ab = sum(r['annual_budget'] for r in rows)
total_cb = sum(r['cumulative_budget'] for r in rows)
total_ce = sum(r['cumulative_executed'] for r in rows)
print(f'   全年预算合计: {total_ab}万')
print(f'   累计预算合计: {total_cb}万')
print(f'   执行合计: {total_ce}万')

# 写入 cockpit.db
conn = sqlite3.connect(DB_PATH)
conn.execute('PRAGMA journal_mode=WAL')

# 按服务中心名称匹配更新
updated = 0
for r in rows:
    # 尝试精确匹配
    result = conn.execute(
        'UPDATE payment_centers SET cumulative_budget = ?, cumulative_executed = ? WHERE center = ?',
        (r['cumulative_budget'], r['cumulative_executed'], r['center'])
    )
    if result.rowcount == 0:
        # 模糊匹配
        result = conn.execute(
            'UPDATE payment_centers SET cumulative_budget = ?, cumulative_executed = ? WHERE center LIKE ?',
            (r['cumulative_budget'], r['cumulative_executed'], f'%{r["center"][-6:]}%')
        )
    updated += result.rowcount

conn.commit()

# 验证
verify = conn.execute(
    'SELECT COUNT(*), SUM(cumulative_budget), SUM(cumulative_executed) FROM payment_centers'
).fetchone()
print(f'\n💾 已更新 {updated} 条')
print(f'   累计预算合计: {verify[1]}万')
print(f'   累计执行合计: {verify[2]}万')
conn.close()

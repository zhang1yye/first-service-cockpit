#!/usr/bin/env python3
"""Import the current active project master workbook into versioned cockpit tables.

Default CLI mode is dry-run. Pass --apply to write only the dedicated
project-profile tables; existing projects/payment/collection tables are never deleted.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SHEET_NAME = "服务面积数据（在管+合约）"
ACTIVE_STATUS = "在管"

SPECIAL_CENTER_LINKS = {
    "第一服务北京满庭青云服务中心": ["第一服务满庭青云服务中心"],
}

# Verified against the workbook's real A:CE coordinates. Tail operational fields
# are intentionally shifted by one column because source row values start in BU,
# while the visible header starts in BT.
COL = {
    "index": 1, "region": 2, "company_entity": 3, "area": 4,
    "service_center": 5, "phase_name": 6, "management_status": 7,
    "planned_movein": 8, "actual_movein": 9, "signed_area": 10,
    "phase_signed_area": 11, "managed_area": 12, "pending_area": 13,
    "area_source": 14, "signed_units": 15, "movedin_units": 16,
    "saturated_income": 17, "avg_property_fee": 18,
    "residential_fee": 19, "villa_fee": 20, "commercial_fee": 21,
    "office_fee": 22, "kindergarten_fee": 23, "club_fee": 24,
    "stacked_villa_fee": 25, "remuneration": 26,
    "property_type": 27, "service_type": 28, "project_source": 29,
    "client_type": 30, "province": 31, "city": 32, "address": 33,
    "city_tier": 34, "signing_entity": 35, "holding_relation": 36,
    "joint_venture": 37, "contract_client": 38, "contract_name": 39,
    "contract_start": 40, "contract_end": 41, "charging_model": 42,
    "exit_time": 43, "exit_reason": 44, "committee": 45,
    "chargeable_area": 46, "residential_charge_area": 47,
    "villa_charge_area": 48, "commercial_charge_area": 49,
    "office_charge_area": 50, "kindergarten_charge_area": 51,
    "club_charge_area": 52, "stacked_villa_charge_area": 53,
    "ground_parking": 54, "underground_parking": 55,
    "civil_defense_parking": 56, "non_motor_parking": 57,
    "ground_parking_fee": 58, "underground_parking_fee": 59,
    "temporary_parking_fee": 60, "parking_execution": 61,
    "contract_filed": 62, "management_room_location": 63,
    "management_room_area": 64, "committee_room_location": 65,
    "committee_room_area": 66, "maintenance_fund_used": 67,
    "maintenance_fund_balance": 68, "vacancy_fee": 69,
    "public_revenue_contract": 70, "public_revenue_accounting": 71,
    # Source-value offset verified by sample rows.
    "green_area": 73, "green_staff": 74, "green_outsourced": 75,
    "entrances": 76, "closed_entrances": 77, "staffed_entrances": 78,
    "unattended_entrances": 79, "operable_area": 80, "cameras": 81,
    "recorders": 82, "freight_elevators": 83, "passenger_elevators": 84,
}

NUMERIC_FIELDS = {
    "signed_area", "phase_signed_area", "managed_area", "pending_area",
    "signed_units", "movedin_units", "saturated_income", "avg_property_fee",
    "residential_fee", "villa_fee", "commercial_fee", "office_fee",
    "kindergarten_fee", "club_fee", "stacked_villa_fee", "remuneration",
    "chargeable_area", "residential_charge_area", "villa_charge_area",
    "commercial_charge_area", "office_charge_area", "kindergarten_charge_area",
    "club_charge_area", "stacked_villa_charge_area", "ground_parking",
    "underground_parking", "civil_defense_parking", "non_motor_parking",
    "ground_parking_fee", "underground_parking_fee", "temporary_parking_fee",
    "management_room_area", "committee_room_area", "maintenance_fund_used",
    "maintenance_fund_balance", "green_area", "green_staff", "entrances",
    "closed_entrances", "staffed_entrances", "unattended_entrances",
    "operable_area", "cameras", "recorders", "freight_elevators",
    "passenger_elevators",
}


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def number(value: Any) -> float | int | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    raw = str(value).strip().replace(",", "").replace("，", "")
    if not raw or raw in {"/", "-", "—", "不涉及", "未入伙", "未交付", "无"}:
        return None
    matches = re.findall(r"-?\d+(?:\.\d+)?", raw)
    if len(matches) != 1:
        return None
    parsed = float(matches[0])
    return int(parsed) if parsed.is_integer() else parsed


def source_headers(ws) -> dict[int, str]:
    headers: dict[int, str] = {}
    used: dict[str, int] = {}
    for col in range(1, 84):
        h1, h2 = text(ws.cell(1, col).value), text(ws.cell(2, col).value)
        label = h2 or h1 or f"未命名列_{col}"
        if label in used:
            used[label] += 1
            label = f"{label}_{used[label]}"
        else:
            used[label] = 1
        headers[col] = label
    return headers


def is_data_index(value: Any) -> bool:
    return (isinstance(value, (int, float)) and not isinstance(value, bool)) or (
        isinstance(value, str) and value.strip().isdigit()
    )


def row_record(ws, row_no: int, headers: dict[int, str], inherited_center: str) -> tuple[dict[str, Any], str]:
    explicit_center = text(ws.cell(row_no, COL["service_center"]).value)
    center = explicit_center or inherited_center
    raw = {headers[c]: json_value(ws.cell(row_no, c).value) for c in range(1, 84)}
    record: dict[str, Any] = {"source_row": row_no, "raw": raw}
    for field, col in COL.items():
        value = ws.cell(row_no, col).value if col <= ws.max_column else None
        if field == "service_center":
            record[field] = center
        elif field in NUMERIC_FIELDS:
            record[field] = number(value)
        elif field in {"planned_movein", "actual_movein", "contract_start", "contract_end", "exit_time"}:
            record[field] = json_value(value)
        else:
            record[field] = text(value)
    return record, center


def _sum(rows: list[dict[str, Any]], field: str) -> float:
    return float(sum(float(row[field]) for row in rows if row.get(field) is not None))


def _first(rows: list[dict[str, Any]], field: str) -> Any:
    return next((row.get(field) for row in rows if row.get(field) not in (None, "")), None)


def parse_active_project_master(workbook_path: str | Path) -> dict[str, Any]:
    workbook_path = Path(workbook_path)
    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"缺少工作表：{SHEET_NAME}")
    ws = wb[SHEET_NAME]
    headers = source_headers(ws)
    phases: list[dict[str, Any]] = []
    inherited_center = ""
    for row_no in range(1, ws.max_row + 1):
        if not is_data_index(ws.cell(row_no, COL["index"]).value):
            continue
        record, inherited_center = row_record(ws, row_no, headers, inherited_center)
        if record["management_status"] != ACTIVE_STATUS:
            continue
        if not record["service_center"] or not record["phase_name"]:
            raise ValueError(f"第{row_no}行在管项目缺少服务中心或项目分期名称")
        phases.append(record)

    grouped: OrderedDict[str, list[dict[str, Any]]] = OrderedDict()
    for phase in phases:
        grouped.setdefault(phase["service_center"], []).append(phase)
    profiles: list[dict[str, Any]] = []
    for center, items in grouped.items():
        profiles.append({
            "service_center": center,
            "region": _first(items, "region") or "华北",
            "area": _first(items, "area") or "",
            "company_entity": _first(items, "company_entity") or "",
            "management_status": ACTIVE_STATUS,
            "phase_count": len(items),
            "property_type": _first(items, "property_type") or "",
            "service_type": _first(items, "service_type") or "",
            "project_source": _first(items, "project_source") or "",
            "client_type": _first(items, "client_type") or "",
            "province": _first(items, "province") or "",
            "city": _first(items, "city") or "",
            "address": _first(items, "address") or "",
            "signed_area": max((float(x["signed_area"]) for x in items if x.get("signed_area") is not None), default=0),
            "phase_signed_area": _sum(items, "phase_signed_area"),
            "managed_area": _sum(items, "managed_area"),
            "pending_area": _sum(items, "pending_area"),
            "signed_units": int(_sum(items, "signed_units")),
            "movedin_units": int(_sum(items, "movedin_units")),
            "source_rows": [x["source_row"] for x in items],
        })
    return {
        "source_file": str(workbook_path),
        "source_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "source_sheet": SHEET_NAME,
        "filter_status": ACTIVE_STATUS,
        "profiles": profiles,
        "phases": phases,
    }


SCHEMA = """
CREATE TABLE IF NOT EXISTS project_profile_import_batches (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  source_file TEXT NOT NULL,
  source_sha256 TEXT NOT NULL UNIQUE,
  source_sheet TEXT NOT NULL,
  filter_status TEXT NOT NULL,
  profile_count INTEGER NOT NULL,
  phase_count INTEGER NOT NULL,
  imported_at TEXT NOT NULL DEFAULT (datetime('now','localtime'))
);
CREATE TABLE IF NOT EXISTS project_profiles (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  batch_id INTEGER NOT NULL,
  service_center TEXT NOT NULL,
  region TEXT NOT NULL DEFAULT '华北',
  area TEXT NOT NULL DEFAULT '',
  company_entity TEXT NOT NULL DEFAULT '',
  management_status TEXT NOT NULL,
  phase_count INTEGER NOT NULL DEFAULT 0,
  property_type TEXT NOT NULL DEFAULT '',
  service_type TEXT NOT NULL DEFAULT '',
  project_source TEXT NOT NULL DEFAULT '',
  client_type TEXT NOT NULL DEFAULT '',
  province TEXT NOT NULL DEFAULT '',
  city TEXT NOT NULL DEFAULT '',
  address TEXT NOT NULL DEFAULT '',
  signed_area REAL NOT NULL DEFAULT 0,
  phase_signed_area REAL NOT NULL DEFAULT 0,
  managed_area REAL NOT NULL DEFAULT 0,
  pending_area REAL NOT NULL DEFAULT 0,
  signed_units INTEGER NOT NULL DEFAULT 0,
  movedin_units INTEGER NOT NULL DEFAULT 0,
  source_rows_json TEXT NOT NULL DEFAULT '[]',
  UNIQUE(batch_id, service_center),
  FOREIGN KEY(batch_id) REFERENCES project_profile_import_batches(id)
);
CREATE TABLE IF NOT EXISTS project_phase_profiles (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  batch_id INTEGER NOT NULL,
  profile_id INTEGER NOT NULL,
  source_row INTEGER NOT NULL,
  phase_name TEXT NOT NULL,
  management_status TEXT NOT NULL,
  payload_json TEXT NOT NULL,
  UNIQUE(batch_id, source_row),
  FOREIGN KEY(batch_id) REFERENCES project_profile_import_batches(id),
  FOREIGN KEY(profile_id) REFERENCES project_profiles(id)
);
CREATE TABLE IF NOT EXISTS project_profile_center_links (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  batch_id INTEGER NOT NULL,
  profile_id INTEGER NOT NULL,
  source_system TEXT NOT NULL,
  source_center TEXT NOT NULL,
  link_method TEXT NOT NULL,
  created_at TEXT NOT NULL DEFAULT (datetime('now','localtime')),
  UNIQUE(batch_id, profile_id, source_system, source_center),
  FOREIGN KEY(batch_id) REFERENCES project_profile_import_batches(id),
  FOREIGN KEY(profile_id) REFERENCES project_profiles(id)
);
CREATE INDEX IF NOT EXISTS idx_project_profiles_batch_area ON project_profiles(batch_id, area, service_center);
CREATE INDEX IF NOT EXISTS idx_project_phase_profiles_profile ON project_phase_profiles(profile_id, source_row);
CREATE INDEX IF NOT EXISTS idx_project_profile_center_links_profile ON project_profile_center_links(profile_id, source_system);
"""


def normalize_center_name(value: Any) -> str:
    return re.sub(
        r"(?:服务中心|体验中心)$",
        "",
        re.sub(r"^第一(?:服务|酒店)", "", text(value).replace("·", "").replace("•", "").replace("・", "").replace(" ", "")),
    )


def normalize_full_center_name(value: Any) -> str:
    return text(value).replace("·", "").replace("•", "").replace("・", "").replace(" ", "")


def table_exists(conn: sqlite3.Connection, table: str) -> bool:
    return conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)
    ).fetchone() is not None


def refresh_center_links(conn: sqlite3.Connection, batch_id: int) -> int:
    conn.execute("DELETE FROM project_profile_center_links WHERE batch_id=?", (batch_id,))
    profiles = conn.execute(
        "SELECT id, service_center FROM project_profiles WHERE batch_id=?", (batch_id,)
    ).fetchall()
    inserted = 0
    for source_system, table in (("payment", "payment_centers"), ("collection", "collection_centers")):
        if not table_exists(conn, table):
            continue
        source_names = [row[0] for row in conn.execute(f"SELECT DISTINCT center FROM {table} WHERE center IS NOT NULL")]
        by_key: dict[str, list[str]] = {}
        by_full_key: dict[str, list[str]] = {}
        for source_name in source_names:
            by_key.setdefault(normalize_center_name(source_name), []).append(source_name)
            by_full_key.setdefault(normalize_full_center_name(source_name), []).append(source_name)
        for profile_id, service_center in profiles:
            targets = SPECIAL_CENTER_LINKS.get(service_center, [service_center])
            method = "explicit_alias" if service_center in SPECIAL_CENTER_LINKS else "normalized_exact"
            for target in targets:
                matches = by_full_key.get(normalize_full_center_name(target), [])
                if len(matches) != 1:
                    candidates = by_key.get(normalize_center_name(target), [])
                    if text(target).endswith("体验中心"):
                        typed = [name for name in candidates if text(name).endswith("体验中心")]
                    else:
                        typed = [name for name in candidates if text(name).endswith("服务中心") and not text(name).endswith("体验中心")]
                    matches = typed if len(typed) == 1 else candidates
                if len(matches) != 1:
                    continue
                conn.execute(
                    """INSERT OR IGNORE INTO project_profile_center_links(
                    batch_id,profile_id,source_system,source_center,link_method
                    ) VALUES (?,?,?,?,?)""",
                    (batch_id, profile_id, source_system, matches[0], method),
                )
                inserted += 1
    return inserted


def import_active_project_master(workbook_path: str | Path, db_path: str | Path) -> dict[str, Any]:
    parsed = parse_active_project_master(workbook_path)
    conn = sqlite3.connect(Path(db_path))
    conn.execute("PRAGMA foreign_keys=ON")
    conn.executescript(SCHEMA)
    existing = conn.execute(
        "SELECT id, profile_count, phase_count FROM project_profile_import_batches WHERE source_sha256=?",
        (parsed["source_sha256"],),
    ).fetchone()
    if existing:
        with conn:
            link_count = refresh_center_links(conn, int(existing[0]))
        conn.close()
        return {
            "status": "already_imported", "batch_id": existing[0],
            "profile_count": existing[1], "phase_count": existing[2],
            "source_sha256": parsed["source_sha256"], "link_count": link_count,
        }
    try:
        with conn:
            cur = conn.execute(
                "INSERT INTO project_profile_import_batches(source_file,source_sha256,source_sheet,filter_status,profile_count,phase_count) VALUES (?,?,?,?,?,?)",
                (parsed["source_file"], parsed["source_sha256"], parsed["source_sheet"], parsed["filter_status"], len(parsed["profiles"]), len(parsed["phases"])),
            )
            batch_id = int(cur.lastrowid)
            profile_ids: dict[str, int] = {}
            for p in parsed["profiles"]:
                cur = conn.execute(
                    """INSERT INTO project_profiles(
                    batch_id,service_center,region,area,company_entity,management_status,phase_count,
                    property_type,service_type,project_source,client_type,province,city,address,
                    signed_area,phase_signed_area,managed_area,pending_area,signed_units,movedin_units,source_rows_json
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (batch_id,p["service_center"],p["region"],p["area"],p["company_entity"],p["management_status"],p["phase_count"],p["property_type"],p["service_type"],p["project_source"],p["client_type"],p["province"],p["city"],p["address"],p["signed_area"],p["phase_signed_area"],p["managed_area"],p["pending_area"],p["signed_units"],p["movedin_units"],json.dumps(p["source_rows"],ensure_ascii=False)),
                )
                profile_ids[p["service_center"]] = int(cur.lastrowid)
            for phase in parsed["phases"]:
                conn.execute(
                    "INSERT INTO project_phase_profiles(batch_id,profile_id,source_row,phase_name,management_status,payload_json) VALUES (?,?,?,?,?,?)",
                    (batch_id,profile_ids[phase["service_center"]],phase["source_row"],phase["phase_name"],phase["management_status"],json.dumps(phase,ensure_ascii=False,default=json_value)),
                )
            link_count = refresh_center_links(conn, batch_id)
    finally:
        conn.close()
    return {
        "status": "imported", "batch_id": batch_id,
        "profile_count": len(parsed["profiles"]), "phase_count": len(parsed["phases"]),
        "source_sha256": parsed["source_sha256"], "link_count": link_count,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--db", type=Path)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--output", type=Path)
    args = ap.parse_args()
    parsed = parse_active_project_master(args.workbook)
    preview = {
        "source_file": parsed["source_file"], "source_sha256": parsed["source_sha256"],
        "source_sheet": parsed["source_sheet"], "filter_status": parsed["filter_status"],
        "profile_count": len(parsed["profiles"]), "phase_count": len(parsed["phases"]),
        "profiles": parsed["profiles"],
    }
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(preview, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.apply:
        if not args.db:
            ap.error("--apply requires --db")
        result = import_active_project_master(args.workbook, args.db)
    else:
        result = {"status": "dry_run", **{k: preview[k] for k in ("source_sha256", "profile_count", "phase_count")}}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
